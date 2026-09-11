import base64
import io
from pathlib import Path

from fastapi.testclient import TestClient
from PIL import Image as PILImage
from pptx import Presentation

from app.brand_engine import flow_deck, render_deck
from app.main import app
from app.schemas import Deck, Image, Page, Table

TEMPLATE = str(
    Path(__file__).parents[1] / "templates" / "2026" / "Salasar_Corporate_Blank.pptx"
)
client = TestClient(app)


def _img_b64() -> str:
    im = PILImage.new("RGB", (400, 300), (0, 112, 192))
    b = io.BytesIO()
    im.save(b, "PNG")
    return base64.b64encode(b.getvalue()).decode()


def test_flow_deck_emits_positioned_placements():
    deck = flow_deck([Page(title="Mix", body="Intro line.",
                           tables=[Table(header=["A", "B"], rows=[["A", "B"], ["1", "2"]])],
                           images=[Image(data=_img_b64())])])
    assert deck.slides
    pls = deck.slides[0].placements
    assert {p.kind for p in pls} == {"body", "table", "image"}
    for p in pls:
        assert p.width > 0 and p.top >= 0        # positioned


def test_render_deck_round_trips_to_valid_pptx():
    deck = flow_deck([Page(title="Title", body="Hello.")])
    data = render_deck(deck, TEMPLATE)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides._sldIdLst) == len(deck.slides)


def test_layout_then_render_endpoints():
    r = client.post("/layout", json={"pages": [{"title": "T", "body": "Hello world."}]})
    assert r.status_code == 200
    deck = r.json()
    assert deck["slides"] and deck["slides"][0]["placements"]

    r2 = client.post("/render", json=deck)
    assert r2.status_code == 200
    assert r2.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    Presentation(io.BytesIO(r2.content))          # opens -> valid


def test_edited_deck_renders_the_edit():
    # edit the body text in the Placement JSON, render, confirm it lands
    deck = flow_deck([Page(title="T", body="Original text.")]).model_dump()
    for p in deck["slides"][0]["placements"]:
        if p["kind"] == "body":
            p["text"] = "Edited body content here."
    data = render_deck(Deck.model_validate(deck), TEMPLATE)
    prs = Presentation(io.BytesIO(data))
    texts = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
    assert any("Edited body content here." in t for t in texts)


def test_render_empty_deck_is_422():
    r = client.post("/render", json={"slides": []})
    assert r.status_code == 422
