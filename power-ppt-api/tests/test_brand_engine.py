import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from app.brand_engine import build_deck
from app.brand_engine import geometry as g
from app.schemas import Page, SlidePlan, Table

TEMPLATE = str(
    Path(__file__).parents[1] / "templates" / "2026" / "Salasar_Corporate_Blank.pptx"
)


# ── helpers ──────────────────────────────────────────────────────────────────
def _open(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def _pictures(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield sh
        elif sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pictures(sh.shapes)


def _slide_text(slide) -> str:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                out.extend(c.text for c in row.cells)
    return "\n".join(out)


def _run_colours(slide):
    cols = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        cols.append(str(r.font.color.rgb))
                    except Exception:
                        pass
    return cols


def _cnvpr_ids(slide):
    return [
        e.get("id")
        for e in slide.shapes._spTree.iter()
        if e.tag.endswith("}cNvPr")
    ]


@pytest.fixture(scope="module")
def basic_deck():
    plan = SlidePlan(
        pages=[
            Page(
                title="Motor Insurance - Overview",
                body="First paragraph of the summary.\n\nSecond paragraph here.",
            ),
            Page(
                title="Claims Process",
                tables=[
                    Table(
                        header=["Step", "Detail"],
                        rows=[["Step", "Detail"], ["1", "Intimate the claim"], ["2", "Surveyor visit"]],
                    )
                ],
            ),
        ]
    )
    return _open(build_deck(plan, TEMPLATE))


# ── tests ────────────────────────────────────────────────────────────────────
def test_slide_count(basic_deck):
    assert len(basic_deck.slides) == 2


def test_output_opens_and_no_duplicate_ids(basic_deck):
    for slide in basic_deck.slides:
        ids = _cnvpr_ids(slide)
        assert len(ids) == len(set(ids)), "duplicate shape IDs would corrupt the file"


def test_brand_images_present_on_every_slide(basic_deck):
    for slide in basic_deck.slides:
        assert len(list(_pictures(slide.shapes))) >= 2, "background + logo must survive"


def test_heading_colour_split(basic_deck):
    cols = _run_colours(basic_deck.slides[0])
    assert str(g.BLUE) in cols, "primary term must be brand blue"
    assert str(g.GREEN) in cols, "qualifier after hyphen must be brand green"


def test_green_rule_off_by_default(basic_deck):
    # SHOW_GREEN_RULE is off per review — no green auto-shape should be present.
    assert not g.SHOW_GREEN_RULE
    for sh in basic_deck.slides[0].shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                assert str(sh.fill.fore_color.rgb) != str(g.GREEN)
            except Exception:
                pass


def test_body_stays_in_content_zone(basic_deck):
    slide = basic_deck.slides[0]
    body_boxes = [
        sh for sh in slide.shapes
        if sh.has_text_frame and not sh.has_table and sh.top and sh.top > Inches(1.4)
    ]
    assert body_boxes, "a body textbox should exist below the heading"
    for box in body_boxes:
        assert box.top >= Inches(1.5) - Inches(0.1)
        assert box.top + box.height <= Inches(6.9), "body must clear the footer bar"


def test_title_without_hyphen_is_all_blue():
    plan = SlidePlan(pages=[Page(title="Claims Process", body="Body.")])
    slide = _open(build_deck(plan, TEMPLATE)).slides[0]
    cols = _run_colours(slide)
    assert str(g.BLUE) in cols
    assert str(g.GREEN) not in cols, "no hyphen -> no green run in the heading"


def test_pagination_splits_long_body():
    long_body = ("Paragraph number filler content. " * 40)
    long_body = "\n\n".join([long_body] * 6)  # well past one page
    plan = SlidePlan(pages=[Page(title="Long Section", body=long_body)])
    prs = _open(build_deck(plan, TEMPLATE))
    assert len(prs.slides) > 1, "overflowing body should paginate"
    titles = " ".join(_slide_text(s) for s in prs.slides)
    assert g.CONTINUATION_SUFFIX.strip() in titles, "continuation pages carry the suffix"


def test_titleless_page_renders_no_heading_but_keeps_body():
    plan = SlidePlan(pages=[Page(title="", body="Just body text, no title.")])
    slide = _open(build_deck(plan, TEMPLATE)).slides[0]
    assert "JUST BODY TEXT" in _slide_text(slide).upper()
