"""
Brand heading rendering (full brand spec, D3):
  top-left, ALL CAPS, primary term Blue, qualifier after the hyphen Green,
  a green rule under the first word, Poppins Semi-Bold.

v1 sizes the text and the green rule with a character-width heuristic so the engine
needs no bundled font file. A precise measurement path (bundled Poppins TTF via
Pillow) can replace `_per_char_in` later without changing this module's interface.
"""

from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from . import geometry as g


def _per_char_in(size_pt: float) -> float:
    """Rough average glyph advance (inches) for ALL-CAPS Poppins Semi-Bold."""
    return (size_pt / 72.0) * 0.62


def _fit_size(text: str, base_pt: int, max_width_in: float) -> int:
    """Shrink the heading only if it would overflow the heading width."""
    if not text:
        return base_pt
    est = len(text) * _per_char_in(base_pt)
    if est <= max_width_in:
        return base_pt
    return max(g.HEAD_MIN_PT, int(base_pt * (max_width_in / est)))


def _split_title(text: str):
    """
    Split on the first hyphen. Primary keeps the hyphen; the remainder (with its
    original spacing) becomes the qualifier. No hyphen -> whole title is primary.
    """
    idx = next((k for k, ch in enumerate(text) if ch in "-–"), -1)
    if idx == -1:
        return text, ""
    return text[: idx + 1], text[idx + 1:]


def render_heading(slide, title: str) -> None:
    if not title or not title.strip():
        return

    text = title.strip().upper()
    size = _fit_size(text, g.HEAD_SIZE_PT, g.HEAD_WIDTH)

    box = slide.shapes.add_textbox(
        Inches(g.HEAD_LEFT), Inches(g.HEAD_TOP),
        Inches(g.HEAD_WIDTH), Inches(g.HEAD_HEIGHT),
    )
    tf = box.text_frame
    tf.word_wrap = False
    para = tf.paragraphs[0]

    primary, qualifier = _split_title(text)

    r1 = para.add_run()
    r1.text = primary
    r1.font.name = g.FONT_HEADING_SB
    r1.font.size = Pt(size)
    r1.font.color.rgb = g.BLUE

    if qualifier:
        r2 = para.add_run()
        r2.text = qualifier
        r2.font.name = g.FONT_HEADING_SB
        r2.font.size = Pt(size)
        r2.font.color.rgb = g.GREEN

    # ── Green rule under the first word (brand spec; toggled off per review) ──
    if g.SHOW_GREEN_RULE:
        first_word = text.split(" ", 1)[0]
        rule_w = max(g.RULE_MIN_W, len(first_word) * _per_char_in(size))
        rule_top = g.HEAD_TOP + (size / 72.0) * 1.25 + g.RULE_GAP

        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(g.HEAD_LEFT), Inches(rule_top),
            Inches(rule_w), Inches(g.RULE_HEIGHT),
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = g.GREEN
        rect.line.fill.background()
        rect.shadow.inherit = False
