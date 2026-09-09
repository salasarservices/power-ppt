"""
Template loading and slide-cloning.

The 2026 template is a flat image canvas: each slide is a group of two pictures
(full-bleed background with the footer bar, plus the logo). It has NO title
placeholder and NO "TITLE GOES HERE" marker — so the engine clones the canvas and
*adds* heading/body shapes onto it. The clone logic (deep-copy the shape tree, copy
relationships, remap rIds, re-number shape ids) is the proven approach salvaged from
the previous implementation.
"""

import copy

from pptx import Presentation
from pptx.oxml.ns import qn

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

_RID_ATTRS = [
    f"{{{_R_NS}}}embed",
    f"{{{_R_NS}}}id",
    f"{{{_R_NS}}}link",
    f"{{{_R_NS}}}href",
]

# Structural header elements of a shape tree — kept, not treated as content.
_STRUCT_TAGS = {qn("p:cNvGrpSpPr"), qn("p:grpSpPr")}


def load_template(path):
    """Open the bundled brand template."""
    return Presentation(path)


def _copy_rels(src_slide, dst_slide):
    """Copy every relationship; return {old_rId: new_rId} for IDs that changed."""
    rId_map = {}
    for rel in list(src_slide.part.rels.values()):
        try:
            new_rId = dst_slide.part.relate_to(
                rel._target, rel.reltype, is_external=rel.is_external
            )
            if new_rId != rel.rId:
                rId_map[rel.rId] = new_rId
        except Exception:
            pass
    return rId_map


def _remap_rids(element, rId_map):
    if not rId_map:
        return
    for el in element.iter():
        for attr in _RID_ATTRS:
            if attr in el.attrib and el.attrib[attr] in rId_map:
                el.attrib[attr] = rId_map[el.attrib[attr]]


def clone_content_slide(prs, source_idx=0):
    """
    Append a full copy of the template slide at `source_idx` (its shapes, images and
    relationships) and return the new slide.
    """
    source = prs.slides[source_idx]
    new_slide = prs.slides.add_slide(source.slide_layout)

    src_tree = source.shapes._spTree
    dst_tree = new_slide.shapes._spTree

    for el in list(dst_tree):
        if el.tag not in _STRUCT_TAGS:
            dst_tree.remove(el)

    copied = [copy.deepcopy(el) for el in src_tree if el.tag not in _STRUCT_TAGS]
    rId_map = _copy_rels(source, new_slide)
    for el in copied:
        _remap_rids(el, rId_map)
        dst_tree.append(el)

    return new_slide


def fix_shape_ids(slide):
    """Re-number every <p:cNvPr id> sequentially so there are no duplicates."""
    counter = 1
    for elem in slide.shapes._spTree.iter():
        if elem.tag.endswith("}cNvPr"):
            elem.set("id", str(counter))
            counter += 1


def remove_slide_number_fields(slide):
    """Drop any shape carrying a <a:fld type="slidenum"> field."""
    sp_tree = slide.shapes._spTree
    to_remove = []
    for sp in list(sp_tree.findall(qn("p:sp"))):
        for fld in sp.findall(f".//{{{_A_NS}}}fld"):
            if "slidenum" in fld.get("type", "").lower():
                to_remove.append(sp)
                break
    for sp in to_remove:
        sp_tree.remove(sp)


def remove_slide(prs, index):
    """Remove the slide at `index` and drop its presentation relationship."""
    xml_slides = prs.slides._sldIdLst
    sld_id_elem = xml_slides[index]
    r_id = sld_id_elem.get(qn("r:id"))
    xml_slides.remove(sld_id_elem)
    if r_id:
        try:
            prs.part.drop_rel(r_id)
        except Exception:
            pass
